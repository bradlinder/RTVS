"""Word (.docx) export formatter for Radio & TV Story Segmenter.

Provides styled Microsoft Word document generation for transcripts, including:
- Arial typography and custom heading layouts
- Speaker labels with bold run formatting
- Timestamp annotations in muted styling
- Highlighted comment callout paragraphs
"""
from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List, Optional

from prs_shared import format_time

try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False


def create_story_docx(
    title: str,
    blocks: List[Dict[str, Any]],
    output_path: Path | str,
    media_name: str = "",
    start_time: float = 0.0,
    end_time: float = 0.0,
    include_speakers: bool = True,
    include_timestamps: bool = True,
    include_comments: bool = True,
    lang_code: str = "en",
    source_segments: Optional[List[Dict[str, Any]]] = None,
) -> bool:
    """Creates a formatted Word (.docx) document from story blocks and metadata."""
    if not DOCX_AVAILABLE:
        return False

    document = Document()
    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(11)

    lang_label = " (Spanish)" if lang_code == "es" else (" (English)" if lang_code == "en" else "")
    document.add_heading(f"{title}{lang_label}", 0)

    if media_name:
        time_range = f" ({format_time(start_time, False)} - {format_time(end_time, False)})" if end_time > 0 else ""
        document.add_paragraph(f"Recording: {media_name}{time_range}")

    last_speaker = None
    for block in blocks:
        speaker = (block.get("speaker") or "").strip() if include_speakers else ""
        p_text = block.get("text", "").strip()
        if not p_text:
            continue

        p = document.add_paragraph()
        if include_timestamps and "start" in block and block["start"] is not None:
            r_time = p.add_run(f"[{format_time(block['start'], False)}] ")
            r_time.font.color.rgb = RGBColor(120, 120, 120)

        is_speaker_change = block.get("is_speaker_change", (speaker != last_speaker))
        if speaker and is_speaker_change and speaker != last_speaker:
            r_spk = p.add_run(f"{speaker}: ")
            r_spk.bold = True
            last_speaker = speaker

        p.add_run(p_text)
        p.paragraph_format.space_after = Pt(6)

        if include_comments:
            src_idx = block.get("_source_index")
            seg_comment = ""
            if src_idx is not None and source_segments and 0 <= src_idx < len(source_segments):
                seg_comment = (source_segments[src_idx].get("comments") or source_segments[src_idx].get("notes", "")).strip()
            elif "comments" in block or "notes" in block:
                seg_comment = str(block.get("comments") or block.get("notes", "")).strip()

            if seg_comment:
                p_note = document.add_paragraph()
                r_note_lbl = p_note.add_run("💬 Comment: ")
                r_note_lbl.bold = True
                r_note_lbl.font.color.rgb = RGBColor(180, 130, 0)
                r_note_txt = p_note.add_run(seg_comment)
                r_note_txt.font.italic = True
                r_note_txt.font.color.rgb = RGBColor(140, 100, 0)
                p_note.paragraph_format.left_indent = Pt(18)
                p_note.paragraph_format.space_after = Pt(6)

    document.save(str(output_path))
    return True
